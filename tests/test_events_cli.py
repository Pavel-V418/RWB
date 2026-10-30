"""Хранилища событий, извлечение текста ТЗ и CLI run / status / report."""

import zlib
from pathlib import Path
from typing import Any

import docx
import pytest
from typer.testing import CliRunner

from specforge import cli
from specforge.core.events import Event, EventBus, EventType, SqliteEventStore
from specforge.llm import FakeLLM
from specforge.orchestrator.registry import Registry
from specforge.tools.docs import DocumentError, extract_text
from tests.test_orchestrator import five_agents


def test_sqlite_event_store_roundtrip_and_bus(tmp_path: Path) -> None:
    store = SqliteEventStore(tmp_path / "e.sqlite")
    bus = EventBus(store)
    seen: list[str] = []
    bus.subscribe(lambda e: seen.append(e.kind))
    bus.emit(Event(kind=EventType.run_started, run_id="r"))
    bus.emit(Event(kind=EventType.stage_passed, run_id="r", stage="s", data={"n": 1}))
    bus.emit(Event(kind=EventType.run_started, run_id="other"))

    events = SqliteEventStore(tmp_path / "e.sqlite").list("r")  # другое соединение
    assert [(s.event.kind, s.event.stage) for s in events] == [
        ("run_started", None),
        ("stage_passed", "s"),
    ]
    assert events[1].event.payload == {"n": 1}
    assert [s.id for s in store.list("r", after_id=events[0].id)] == [events[1].id]
    assert seen == ["run_started", "stage_passed", "run_started"]


def test_bus_survives_broken_store() -> None:
    class Broken:
        def append(self, event: Event) -> int:
            raise OSError("db down")

        def list(self, run_id: str, *, after_id: int = 0, limit: int = 1000) -> list[Any]:
            return []

    seen: list[Event] = []
    bus = EventBus(Broken())
    bus.subscribe(seen.append)
    bus.emit(Event(kind="x"))
    bus.emit(Event(kind="y"))
    assert [e.kind for e in seen] == ["x", "y"]


def _pdf(text: str) -> bytes:
    """Минимальный PDF с одной страницей текста."""
    stream = zlib.compress(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1"))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d /Filter /FlateDecode >>stream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def test_extract_text_pdf_docx_and_errors(tmp_path: Path) -> None:
    pdf = tmp_path / "req.pdf"
    pdf.write_bytes(_pdf("Requirements from PDF"))
    assert "Requirements from PDF" in extract_text(pdf)

    path = tmp_path / "req.docx"
    document = docx.Document()
    document.add_paragraph("Требование 1")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "поле"
    table.rows[0].cells[1].text = "значение"
    document.save(str(path))
    assert extract_text(path) == "Требование 1\n\nполе | значение"

    with pytest.raises(DocumentError, match="не поддерживается"):
        extract_text(tmp_path / "x.rtf")
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    with pytest.raises(DocumentError, match="DOCX"):
        extract_text(broken)


@pytest.fixture
def cli_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv("SPECFORGE_PG_DSN", "")
    monkeypatch.setenv("SPECFORGE_RUNS_DIR", str(tmp_path / "runs"))
    monkeypatch.setenv("SPECFORGE_STAGE_CACHE", "false")
    agents = five_agents()
    monkeypatch.setattr(cli, "build_registry", lambda: Registry(agents.values()))
    monkeypatch.setattr(cli, "build_llm", lambda: FakeLLM([]))
    requirements = tmp_path / "req.docx"
    document = docx.Document()
    document.add_paragraph("ТЗ из docx")
    document.save(str(requirements))
    return requirements


def test_cli_run_status_report(cli_env: Path) -> None:
    runner = CliRunner()
    result = runner.invoke(cli.app, ["run", str(cli_env), "--run-id", "r1", "--only", "b"])
    assert result.exit_code == 0, result.output
    assert "Итог: finished" in result.output
    assert "passed" in result.output

    runs = Path(cli_env.parent / "runs" / "r1" / "artifacts" / "input")
    assert (runs / "requirements.md").read_text("utf-8") == "ТЗ из docx"

    status = runner.invoke(cli.app, ["status", "r1"])
    assert status.exit_code == 0, status.output
    assert "r1" in status.output
    assert "skipped" in status.output  # c и d не выбраны в --only

    report = runner.invoke(cli.app, ["report", "r1"])
    assert report.exit_code == 0, report.output
    md = cli_env.parent / "runs" / "r1" / "report" / "report.md"
    assert str(md.name) in report.output
    text = md.read_text("utf-8")
    assert "# Отчёт о прогоне `r1`" in text
    assert "| b | ✅ passed | 1 |" in text
    assert (md.parent / "stages_gantt.png").is_file()


def test_cli_run_approve_after_and_conflicting_flags(cli_env: Path) -> None:
    runner = CliRunner()
    paused = runner.invoke(cli.app, ["run", str(cli_env), "--run-id", "r2", "--approve-after", "a"])
    assert paused.exit_code == 0, paused.output
    assert "Ждёт подтверждения после a: specforge resume r2" in paused.output

    both = runner.invoke(cli.app, ["run", str(cli_env), "--autopilot", "--approve-after", "a"])
    assert both.exit_code == 2
    assert runner.invoke(cli.app, ["status", "nope"]).exit_code == 2
