"""Текст требований из .md / .txt / .pdf / .docx — детерминированно, без LLM."""

from pathlib import Path

SUPPORTED_SUFFIXES = (".md", ".txt", ".pdf", ".docx")


class DocumentError(ValueError):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        return _pdf_text(path)
    if suffix == ".docx":
        return _docx_text(path)
    raise DocumentError(
        f"{path.name}: формат {suffix or '(без расширения)'} не поддерживается; "
        f"ожидается {', '.join(SUPPORTED_SUFFIXES)}"
    )


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader  # noqa: PLC0415 — тяжёлый импорт только для pdf
    from pypdf.errors import PdfReadError  # noqa: PLC0415

    try:
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, OSError) as exc:
        raise DocumentError(f"{path.name}: не читается как PDF: {exc}") from exc
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _docx_text(path: Path) -> str:
    import docx  # noqa: PLC0415 — тяжёлый импорт только для docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx бросает разные ошибки для битых файлов
        raise DocumentError(f"{path.name}: не читается как DOCX: {exc}") from exc
    blocks: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)
