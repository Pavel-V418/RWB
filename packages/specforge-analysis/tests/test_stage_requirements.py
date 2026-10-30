import json
from pathlib import Path

import pytest

from specforge_analysis import stage_requirements
from specforge_analysis.assemble import build_base, merge_analysis
from specforge_analysis.gates import requirements_gate as gate
from specforge_analysis.llm import FakeLLM, LLMSettings, Metrics, OpenAICompatLLM
from specforge_analysis.normalize import SourceIndex, load_requirements
from specforge_analysis.schemas import AAmbiguity, Analysis, Extraction

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "cases" / "case1_parking"
FIX = ROOT / "golden" / "case1_parking" / "llm"


def _run(tmp_path, **kw):
    llm = FakeLLM(FIX, metrics=Metrics())
    return stage_requirements.run(CASE, tmp_path, llm, **kw), tmp_path


def test_golden_pipeline_passes(tmp_path):
    res, out = _run(tmp_path)
    assert res.passed, [str(e) for e in res.errors]
    req = json.loads((out / "requirements.json").read_text(encoding="utf-8"))
    assert len(req["functional"]) == 21
    assert all(f["acceptance"] for f in req["functional"])
    assert len(req["questions"]) == 5
    assert len(list((out / "acceptance").glob("*.feature"))) == 21
    assert res.llm_calls == 4  # extract, analyze, critic, repair


def test_ids_are_stable(tmp_path):
    _, a = _run(tmp_path / "a")
    _, b = _run(tmp_path / "b")
    assert (a / "requirements.json").read_text(encoding="utf-8") == (b / "requirements.json").read_text(
        encoding="utf-8"
    )


def test_customer_items_have_source_lines(tmp_path):
    _, out = _run(tmp_path)
    req = json.loads((out / "requirements.json").read_text(encoding="utf-8"))
    for f in req["functional"]:
        if f["origin"] == "customer":
            assert f["source"]["line"], f["id"]


def test_interactive_mode_waits_for_answers(tmp_path):
    res, _ = _run(tmp_path, mode="interactive")
    assert not res.passed
    assert stage_requirements.needs_answers(res)


def test_answers_are_applied(tmp_path):
    inp = tmp_path / "in"
    inp.mkdir()
    (inp / "requirements.md").write_text((CASE / "requirements.md").read_text(encoding="utf-8"), encoding="utf-8")
    (inp / "answers.yaml").write_text("Q-01: Корпоративный SSO по OpenID Connect\n", encoding="utf-8")
    llm = FakeLLM(FIX, metrics=Metrics())
    res = stage_requirements.run(inp, tmp_path / "out", llm, mode="interactive")
    req = json.loads((tmp_path / "out" / "requirements.json").read_text(encoding="utf-8"))
    q1 = req["questions"][0]
    assert q1["status"] == "answered" and q1["origin"] == "clarification"
    assert any(e.code == "E_Q_OPEN" and e.target == "Q-02" for e in res.errors)


# --------------------------------------------------------------------------- шлюз и политика


def _base():
    text = load_requirements(CASE / "requirements.md")
    ex = Extraction.model_validate_json((FIX / "extract.json").read_text(encoding="utf-8"))
    return build_base(ex, SourceIndex(text), "requirements.md", "auto")


def test_hallucinated_quote_is_downgraded():
    ex = Extraction.model_validate_json((FIX / "extract.json").read_text(encoding="utf-8"))
    ex.functional[0].source_quote = "система должна летать на Луну"
    text = load_requirements(CASE / "requirements.md")
    req = build_base(ex, SourceIndex(text), "requirements.md", "auto")
    fixes = gate.autofix(req)
    downgraded = [f for f in req.functional if f.statement == ex.functional[0].statement][0]
    assert downgraded.origin == "derived"
    assert any(i.code == "W_SOURCE_DOWNGRADED" for i in fixes)


def test_missing_acceptance_is_error():
    req = _base()
    an = Analysis.model_validate_json((FIX / "analyze.json").read_text(encoding="utf-8"))
    an.acceptance = [a for a in an.acceptance if a.fr_id != "FR-4"]
    merge_analysis(req, an, "auto")
    assert any(i.code == "E_NO_AC" and i.target == "FR-4" for i in gate.check(req))


def test_question_limit_and_overflow_to_assumptions():
    req = _base()
    an = Analysis.model_validate_json((FIX / "analyze.json").read_text(encoding="utf-8"))
    extra = [AAmbiguity(topic=f"T{i}", description="d", impact="architecture", severity="high",
                        options=["a", "b"], default_option="a", default_is_safe=False) for i in range(4)]
    an.ambiguities = an.ambiguities + extra
    merge_analysis(req, an, "auto")
    assert len(req.questions) == 5
    assert any(a.statement.startswith("T") for a in req.assumptions)


def test_quote_locate_exact_fuzzy_missing():
    idx = SourceIndex("строка один\nСистема должна:\n- определять местоположение автомобиля;\n")
    assert idx.locate("определять местоположение автомобиля").line == 3
    m = idx.locate("определять текущее местоположение автомобиля")
    assert m and not m.exact
    assert idx.locate("совсем другой текст про погоду") is None


# --------------------------------------------------------------------------- клиент LLM


class _Resp:
    def __init__(self, content, status=200):
        self.status_code = status
        self._content = content
        self.text = content

    def json(self):
        return {"choices": [{"message": {"content": self._content}}],
                "usage": {"prompt_tokens": 100, "completion_tokens": 50}}


def test_llm_retry_on_invalid_json_and_cache(tmp_path, monkeypatch):
    from specforge_analysis import llm as llm_mod
    from specforge_analysis.schemas import Critique

    calls = []
    answers = iter(["<think>...</think> не JSON", '{"findings": []}'])

    def fake_post(url, json=None, headers=None, timeout=None):
        calls.append(json)
        return _Resp(next(answers))

    monkeypatch.setattr(llm_mod.requests, "post", fake_post)
    s = LLMSettings(cache_dir=str(tmp_path / "cache"))
    client = OpenAICompatLLM(s, metrics=Metrics())
    msgs = [{"role": "system", "content": "s"}, {"role": "user", "content": "u"}]
    out = client.structured("critic", msgs, Critique, thinking=False)
    assert out.findings == []
    assert len(calls) == 2
    assert calls[0]["chat_template_kwargs"] == {"enable_thinking": False}
    assert calls[0]["response_format"]["type"] == "json_schema"
    assert [c.attempt for c in client.metrics.calls] == [1, 2]


def test_llm_cache_hit_costs_no_tokens(tmp_path, monkeypatch):
    from specforge_analysis import llm as llm_mod
    from specforge_analysis.schemas import Critique

    monkeypatch.setattr(llm_mod.requests, "post", lambda *a, **k: _Resp('{"findings": []}'))
    s = LLMSettings(cache_dir=str(tmp_path / "cache"))
    msgs = [{"role": "system", "content": "s"}, {"role": "user", "content": "u"}]
    OpenAICompatLLM(s, metrics=Metrics()).structured("critic", msgs, Critique)

    monkeypatch.setattr(llm_mod.requests, "post", lambda *a, **k: pytest.fail("ожидался кэш"))
    client = OpenAICompatLLM(s, metrics=Metrics())
    client.structured("critic", msgs, Critique)
    rec = client.metrics.calls[0]
    assert rec.cached and rec.input_tokens == 0 and rec.output_tokens == 0


def test_docx_input(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Требования", level=1)
    d.add_paragraph("Система должна вести учёт заявок.")
    p = tmp_path / "requirements.docx"
    d.save(p)
    text = load_requirements(p)
    assert "# Требования" in text and "учёт заявок" in text


def test_business_requirements_traced(tmp_path):
    res, out = _run(tmp_path)
    req = json.loads((out / "requirements.json").read_text(encoding="utf-8"))
    assert req["functional"][0]["id"] == "FR-1"
    brs = {b["id"]: b for b in req["business"]}
    assert brs and all(b["line"] > 0 for b in brs.values())
    traced = {fr for b in brs.values() for fr in b["fr"]}
    customer = {f["id"] for f in req["functional"] if f["origin"] == "customer"}
    assert customer <= traced
    assert req["scenarios"][0]["id"] == "UC_001"


def test_llm_waits_on_rate_limit(tmp_path, monkeypatch):
    from specforge_analysis import llm as llm_mod
    from specforge_analysis.schemas import Critique

    answers = iter([_Resp('{"error": "Limit resets at: 2000-01-01 00:00:00 UTC"}', 429), _Resp('{"findings": []}')])
    sent = []
    monkeypatch.setattr(llm_mod.requests, "post", lambda *a, **k: (sent.append(k["json"]), next(answers))[1])
    monkeypatch.setattr(llm_mod.time, "sleep", lambda s: None)
    s = LLMSettings(cache_dir="", tpm_limit=30000)
    msgs = [{"role": "system", "content": "s"}, {"role": "user", "content": "u"}]
    out = OpenAICompatLLM(s, metrics=Metrics()).structured("critic", msgs, Critique, max_tokens=24000)
    assert out.findings == [] and len(sent) == 2
    assert sent[0]["max_tokens"] <= 30000 - 500


def test_system_actor_becomes_role(tmp_path):
    from specforge_analysis.stage_requirements import run
    req = _base()
    an = Analysis.model_validate_json((FIX / "analyze.json").read_text(encoding="utf-8"))
    merge_analysis(req, an, "auto")
    req.scenarios[0].actor = "system"
    req.stories[0].role = req.roles[0].key
    fixes = gate.autofix(req)
    assert req.scenarios[0].actor == "ROLE-system" and req.stories[0].role == req.roles[0].id
    assert any(i.code == "W_ROLE_ADDED" for i in fixes)
    assert not [e for e in gate.check(req) if e.code == "E_REF"]
