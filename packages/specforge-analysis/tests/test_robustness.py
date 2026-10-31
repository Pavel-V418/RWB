"""Устойчивость конвейера: ошибки, которые раньше останавливали этап без пути восстановления."""

import json
import shutil
from pathlib import Path

import pytest

from specforge_analysis import stage_requirements, stage_system
from specforge_analysis.gates import system_gate as gate
from specforge_analysis.gates.pgcheck import check_ddl
from specforge_analysis.llm import FakeLLM, LLMSettings, Metrics, OpenAICompatLLM
from specforge_analysis.render import contracts
from specforge_analysis.schemas import Requirements
from specforge_analysis.schemas_system import ApiDesign, DomainModel, DField

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "cases" / "case1_parking"
FIX = ROOT / "golden" / "case1_parking" / "llm"


@pytest.fixture(scope="module")
def stage1(tmp_path_factory):
    out = tmp_path_factory.mktemp("s1")
    stage_requirements.run(CASE, out, FakeLLM(FIX, metrics=Metrics()))
    return out


def _model(stage1) -> contracts.Model:
    req = Requirements.model_validate_json((stage1 / "requirements.json").read_text(encoding="utf-8"))
    dm = DomainModel.model_validate_json((FIX / "domain.json").read_text(encoding="utf-8"))
    ad = ApiDesign.model_validate_json((FIX / "api_design.json").read_text(encoding="utf-8"))
    return contracts.Model(dm, ad, req)


def _fixtures(tmp_path, edit=None) -> Path:
    fx = tmp_path / "fx"
    shutil.copytree(FIX, fx)
    if edit:
        edit(fx)
    return fx


# ------------------------------------------------------------------ DDL: зарезервированные слова


def test_reserved_column_is_quoted(stage1):
    m = _model(stage1)
    m.entities["Vehicle"].fields.append(DField(name="order", type="int", required=False, description="порядок"))
    m.entities["Vehicle"].fields.append(DField(name="from", type="datetime", required=False))
    ddl = contracts.build_ddl(m)
    assert '    "order" integer' in ddl and '    "from" timestamptz' in ddl
    assert 'COMMENT ON COLUMN vehicles."order"' in ddl
    issues = gate.check(m, contracts.build_openapi(m), None, None)
    assert {i.target for i in issues if i.code == "W_DDL_RESERVED_NAME"} == {"Vehicle"}
    res = check_ddl(ddl)
    if res.ran:
        assert res.ok, res.message


def test_reserved_entity_table_and_enum_are_quoted(stage1):
    m = _model(stage1)
    from specforge_analysis.schemas_system import DEntity
    m.domain.entities.append(DEntity(name="User", description="u", fields=[
        DField(name="end", type="enum", enum_values=["a", "b"])]))
    m = contracts.Model(m.domain, m.design, m.req)
    assert contracts.table_name("User") == "users"          # «users» не зарезервировано
    ddl = contracts.build_ddl(m)
    assert '"end" user_end' in ddl
    res = check_ddl(ddl)
    if res.ran:
        assert res.ok, res.message


# ------------------------------------------------------------------ пути: глаголы


@pytest.mark.parametrize("verb,expected", [("stop", "/stops"), ("start", "/starts"), ("run", "/runs"),
                                           ("send", "/sendings"), ("mark", "/marks")])
def test_action_verbs_autofixed(stage1, verb, expected):
    m = _model(stage1)
    op = next(o for o in m.design.operations if o.operation_id == "updateSession")
    op.method, op.path = "post", f"/api/v1/sessions/{{sessionId}}/{verb}"
    fixes = gate.autofix(m)
    assert op.path == "/api/v1/sessions/{sessionId}" + expected
    assert any(i.code == "W_PATH_RENAMED" for i in fixes)
    assert not any(i.code == "E_PATH_VERB" for i in gate.check(m, contracts.build_openapi(m), None, None))


def test_redundant_verb_dropped_only_with_matching_method(stage1):
    m = _model(stage1)
    ops = [o for o in m.design.operations if o.method == "post"]
    a, b = ops[0], ops[1]
    a.path = "/api/v1/widgets/create"
    b.method, b.path = "get", "/api/v1/gadgets/create"          # GET + create — не убирается молча
    gate.autofix(m)
    assert a.path == "/api/v1/widgets"
    assert b.path == "/api/v1/gadgets/create"
    assert any(i.code == "E_PATH_VERB" for i in gate.check(m, contracts.build_openapi(m), None, None))


# ------------------------------------------------------------------ роли


def test_role_aliases_instead_of_drop(stage1):
    m = _model(stage1)
    keys = [r.key for r in m.req.roles]
    target = keys[0]
    op = m.design.operations[0]
    op.roles = [f"ROLE-{target}", target.upper() + "S"]
    gate.autofix(m)
    assert op.roles == [target]


def test_unknown_roles_fall_back_to_fr_actors(stage1):
    m = _model(stage1)
    op = next(o for o in m.design.operations if o.fr and any(
        f.actors for f in m.req.functional if f.id in o.fr))
    op.roles = ["nobody_knows"]
    fixes = gate.autofix(m)
    assert op.roles, "операция не должна остаться без ролей"
    assert any(i.code == "W_ROLE_FROM_FR" for i in fixes)


# ------------------------------------------------------------------ сбой LLM


def test_llm_failure_returns_gate_result_and_clears_stale(tmp_path):
    out = tmp_path / "out"
    res = stage_requirements.run(CASE, out, FakeLLM(FIX, metrics=Metrics()))
    assert res.passed and (out / "requirements.json").exists()

    broken = tmp_path / "broken"
    broken.mkdir()
    shutil.copy(FIX / "extract.json", broken / "extract.json")          # analyze отсутствует
    res = stage_requirements.run(CASE, out, FakeLLM(broken, metrics=Metrics()))
    assert not res.passed and res.errors[0].code == "E_LLM"
    assert not (out / "requirements.json").exists(), "устаревший результат не должен остаться"
    assert not list((out / "acceptance").glob("*.feature"))
    rep = json.loads((out / "gate_result.json").read_text(encoding="utf-8"))
    assert rep["passed"] is False and rep["errors"][0]["code"] == "E_LLM"


def test_stage2_llm_failure(tmp_path, stage1):
    broken = tmp_path / "broken"
    broken.mkdir()
    shutil.copy(FIX / "domain.json", broken / "domain.json")             # api_design отсутствует
    res = stage_system.run(stage1, tmp_path / "s2", FakeLLM(broken, metrics=Metrics()), check_db=False)
    assert not res.passed and res.errors[0].code == "E_LLM"


def test_output_equal_to_input_is_not_cleaned(tmp_path):
    case = tmp_path / "case"
    shutil.copytree(CASE, case)
    stage_requirements.run(case, case, FakeLLM(FIX, metrics=Metrics()))
    assert (case / "requirements.md").exists() and (case / "expected.yaml").exists()


# ------------------------------------------------------------------ лимит длины ответа


class _Resp:
    def __init__(self, content, finish="stop"):
        self.status_code, self.text = 200, content
        self._d = {"choices": [{"message": {"content": content}, "finish_reason": finish}],
                   "usage": {"prompt_tokens": 10, "completion_tokens": 5}}

    def json(self):
        return self._d


def test_length_truncation_retries_with_bigger_budget(tmp_path, monkeypatch):
    from specforge_analysis import llm as llm_mod
    from specforge_analysis.schemas import Critique

    bodies = []
    answers = iter([_Resp('{"findings": [', "length"), _Resp('{"findings": [', "length"),
                    _Resp('{"findings": []}')])

    def fake_post(url, json=None, headers=None, timeout=None):
        bodies.append(json)
        return next(answers)

    monkeypatch.setattr(llm_mod.requests, "post", fake_post)
    client = OpenAICompatLLM(LLMSettings(cache_dir=""), metrics=Metrics())
    msgs = [{"role": "system", "content": "s"}, {"role": "user", "content": "u"}]
    out = client.structured("critic", msgs, Critique, thinking=True, max_tokens=1000)
    assert out.findings == []
    assert [b["max_tokens"] for b in bodies] == [1000, 1500, 2250]
    assert [b["chat_template_kwargs"]["enable_thinking"] for b in bodies] == [True, True, False]
    assert len(bodies[2]["messages"]) == 2, "обрезанный ответ не добавляется в диалог"
    assert client.metrics.summary()["retries"] == 2


# ------------------------------------------------------------------ ремонт ошибок шлюза этапа 2


def test_design_repair_fixes_blocking_error(tmp_path, stage1):
    def edit(fx):
        d = json.loads((fx / "api_design.json").read_text(encoding="utf-8"))
        op = next(o for o in d["operations"] if o["operation_id"] == "updateSession")
        good = dict(op)
        op["roles"] = []                                    # E_OP_NO_ROLE после autofix…
        op["fr"] = []                                       # …и ролям неоткуда взяться из FR
        (fx / "api_design.json").write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        (fx / "repair_design_1.json").write_text(json.dumps({"operations": [good]}, ensure_ascii=False),
                                                 encoding="utf-8")

    fx = _fixtures(tmp_path, edit)
    res = stage_system.run(stage1, tmp_path / "s2", FakeLLM(fx, metrics=Metrics()), check_db=False)
    assert res.passed, [str(e) for e in res.errors]
    assert any(w.code == "W_DESIGN_REPAIRED" for w in res.warnings)
    assert res.llm_calls == 4  # domain, api_design, repair_design_1, sa_prose


# ------------------------------------------------------------------ этап B: честность трассировки


def test_heuristic_br_links_are_candidates_not_coverage(tmp_path):
    res = stage_requirements.run(CASE, tmp_path, FakeLLM(FIX, metrics=Metrics()))
    req = Requirements.model_validate_json((tmp_path / "requirements.json").read_text(encoding="utf-8"))
    br36 = next(b for b in req.business if b.line == 36)          # «До окончания … 10 минут»
    assert br36.fr == [] and br36.fr_candidates                  # подобран кодом — только кандидат
    w = next(w for w in res.warnings if w.code == "W_BR_UNCOVERED")
    assert "36 (кандидат" in w.message
    assert "(кандидат)" in (tmp_path / "requirements.md").read_text(encoding="utf-8")
    # в трассировке FR кандидат не засчитан
    import csv
    rows = list(csv.DictReader((tmp_path / "traceability.csv").open(encoding="utf-8")))
    assert all(r["business_req"] != br36.id for r in rows)


def test_altered_number_in_quote_is_not_customer():
    from specforge_analysis.normalize import SourceIndex
    text = "Заголовок\n\n- «До окончания парковки осталось 10 минут».\n- «Вы покинули зону».\n"
    ix = SourceIndex(text)
    assert ix.locate("До окончания парковки осталось 30 минут") is None
    m = ix.locate("Вы покинули текущую зону")
    assert m is None or m.line == 4


def test_fuzzy_match_points_to_best_line():
    from specforge_analysis.normalize import SourceIndex
    ix = SourceIndex("первая строка про погоду\nвторая про другое\nсистема должна хранить историю сессий водителя\n")
    m = ix.locate("Система должна хранить историю всех сессий")
    assert m and not m.exact and m.line == 3


def test_critic_blocker_surfaces_in_gate(tmp_path):
    def edit(fx):
        (fx / "critic.json").write_text(json.dumps({"findings": [
            {"severity": "blocker", "category": "missing_flow", "target": "GLOBAL",
             "issue": "Не покрыта выгрузка отчёта.", "fix": "Добавить FR."}]}, ensure_ascii=False), encoding="utf-8")
        (fx / "repair_acceptance_1.json").unlink()
    fx = _fixtures(tmp_path, edit)
    res = stage_requirements.run(CASE, tmp_path / "out", FakeLLM(fx, metrics=Metrics()))
    assert any(w.code == "W_CRITIC_BLOCKER" for w in res.warnings)


# ------------------------------------------------------------------ этап C


def test_interactive_empty_answer_means_default(tmp_path):
    case = tmp_path / "case"
    shutil.copytree(CASE, case)
    first = stage_requirements.run(case, tmp_path / "o1", FakeLLM(FIX, metrics=Metrics()), mode="interactive")
    assert any(e.code == "E_Q_OPEN" for e in first.errors)
    shutil.copy(tmp_path / "o1" / "answers.template.yaml", case / "answers.yaml")   # все ответы пустые
    res = stage_requirements.run(case, tmp_path / "o2", FakeLLM(FIX, metrics=Metrics()), mode="interactive")
    assert res.passed, [str(e) for e in res.errors]
    req = Requirements.model_validate_json((tmp_path / "o2" / "requirements.json").read_text(encoding="utf-8"))
    assert {q.status for q in req.questions} == {"defaulted"}


def test_answer_triggers_acceptance_rebuild(tmp_path):
    case = tmp_path / "case"
    shutil.copytree(CASE, case)
    stage_requirements.run(case, tmp_path / "o1", FakeLLM(FIX, metrics=Metrics()))
    req = Requirements.model_validate_json((tmp_path / "o1" / "requirements.json").read_text(encoding="utf-8"))
    q = next(q for q in req.questions if any(a.startswith("FR-") for a in q.affected))
    (case / "answers.yaml").write_text(f"{q.id}: Решение заказчика XYZ\n", encoding="utf-8")

    seen = []

    class Spy(FakeLLM):
        def structured(self, name, messages, schema, **kw):
            if name.startswith("repair_acceptance"):
                seen.append(messages[-1]["content"])
            return super().structured(name, messages, schema, **kw)

    stage_requirements.run(case, tmp_path / "o2", Spy(FIX, metrics=Metrics()), mode="interactive")
    assert seen and "Решение заказчика XYZ" in seen[0]


def test_docx_tables_keep_position(tmp_path):
    import docx
    from specforge_analysis.normalize import load_requirements
    d = docx.Document()
    d.add_heading("Роли", level=2)
    d.add_paragraph("Права ролей:")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Роль", "Права"
    t.cell(1, 0).text, t.cell(1, 1).text = "Кладовщик", "Выдача"
    d.add_heading("Интеграции", level=2)
    d.add_paragraph("Система учёта ТМЦ.")
    p = tmp_path / "r.docx"
    d.save(p)
    text = load_requirements(p)
    assert text.index("| Кладовщик | Выдача |") < text.index("## Интеграции")
    assert text.index("Права ролей:") < text.index("| Роль | Права |")


def test_rehearse_exit_codes():
    from specforge_analysis.rehearse import CaseScore, exit_code
    ok = CaseScore("a", stage1="пройден", stage2="пройден", docs="пройден")
    skip = CaseScore("b", stage1="нужна модель", stage2="нужна модель", docs="нужна модель")
    bad_docs = CaseScore("c", stage1="пройден", stage2="пройден", docs="ошибок 1 (E_DOC_LINK)")
    assert exit_code([ok, skip]) == 0
    assert exit_code([skip]) == 2
    assert exit_code([ok, bad_docs]) == 1


def test_rehearse_reports_llm_failure(tmp_path):
    from specforge_analysis import rehearse
    root = tmp_path / "r"
    cases, golden = root / "cases", root / "golden" / "caseX" / "llm"
    shutil.copytree(CASE, cases / "caseX")
    golden.mkdir(parents=True)
    shutil.copy(FIX / "extract.json", golden / "extract.json")       # analyze отсутствует
    [s] = rehearse.run(cases, tmp_path / "runs", "fake", check_db=False)
    assert "E_LLM" in s.stage1 and s.error and "analyze" in s.error


def test_non_ecma_regex_is_error(stage1):
    m = _model(stage1)
    m.entities["Vehicle"].fields[1].pattern = r"(?i)^[a-z]{3}$"
    issues = gate.check(m, contracts.build_openapi(m), None, None)
    assert any(i.code == "E_REGEX" and "ECMA" in i.message for i in issues)
